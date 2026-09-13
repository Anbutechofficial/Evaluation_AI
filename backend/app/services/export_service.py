import os
import logging
from typing import List, Dict, Any
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.config import settings

logger = logging.getLogger(__name__)

class ExportService:
    @staticmethod
    def generate_pdf_mark_statement(exam_name: str, subject: str, class_name: str, submissions: List[Dict[str, Any]]) -> str:
        """
        Generates a PDF mark statement using ReportLab.
        """
        filename = f"mark_statement_{exam_name.replace(' ', '_')[:25]}.pdf"
        output_path = os.path.join(settings.EXPORT_DIR, filename)

        doc = SimpleDocTemplate(
            output_path,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        # Custom Styles
        title_style = ParagraphStyle(
            'EduEvalTitle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1e1b4b'),
            alignment=1, # Center
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            'EduEvalSubtitle',
            parent=styles['Normal'],
            fontSize=11,
            leading=14,
            textColor=colors.HexColor('#475569'),
            alignment=1,
            spaceAfter=15
        )
        header_cell_style = ParagraphStyle(
            'HeaderCell',
            parent=styles['Normal'],
            fontSize=10,
            leading=12,
            textColor=colors.white,
            fontName='Helvetica-Bold',
            alignment=1
        )
        body_cell_style = ParagraphStyle(
            'BodyCell',
            parent=styles['Normal'],
            fontSize=9,
            leading=11,
            textColor=colors.HexColor('#1e293b'),
            alignment=1
        )

        elements.append(Paragraph(f"Evaluation AI &bull; Official Mark Statement", title_style))
        elements.append(Paragraph(f"<b>Exam:</b> {exam_name} &nbsp;|&nbsp; <b>Subject:</b> {subject} &nbsp;|&nbsp; <b>Class/Dept:</b> {class_name}", subtitle_style))

        # Table Header
        headers = ["Rank", "Register No", "Student Name", "Score", "Max Marks", "Percentage", "Status"]
        table_data = [[Paragraph(h, header_cell_style) for h in headers]]

        for idx, sub in enumerate(submissions, 1):
            rank_str = str(sub.get("rank") or idx)
            reg_no = sub.get("register_number", "-")
            name = sub.get("student_name", "-")
            score = f"{sub.get('total_score', 0):.1f}"
            max_m = f"{sub.get('max_score', 20):.1f}"
            pct = f"{sub.get('percentage', 0):.1f}%"
            status_str = sub.get("status", "SUBMITTED").replace("_", " ")

            row = [
                Paragraph(rank_str, body_cell_style),
                Paragraph(reg_no, body_cell_style),
                Paragraph(name, body_cell_style),
                Paragraph(score, body_cell_style),
                Paragraph(max_m, body_cell_style),
                Paragraph(pct, body_cell_style),
                Paragraph(status_str, body_cell_style),
            ]
            table_data.append(row)

        col_widths = [50, 90, 180, 70, 70, 80, 110]
        table = Table(table_data, colWidths=col_widths, repeatRows=1)
        
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f46e5')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')])
        ]))

        elements.append(table)
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("<i>Generated automatically by Evaluation AI Platform</i>", ParagraphStyle('Footer', parent=styles['Italic'], fontSize=8, textColor=colors.HexColor('#94a3b8'), alignment=1)))

        doc.build(elements)
        return output_path

    @staticmethod
    def generate_excel_mark_statement(exam_name: str, subject: str, class_name: str, submissions: List[Dict[str, Any]]) -> str:
        """
        Generates an Excel mark statement (.xlsx) with formatting.
        """
        filename = f"mark_statement_{exam_name.replace(' ', '_')[:25]}.xlsx"
        output_path = os.path.join(settings.EXPORT_DIR, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mark Statement"

        # Styles
        title_font = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        border_side = Side(border_style="thin", color="CBD5E1")
        border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center_align = Alignment(horizontal="center", vertical="center")

        # Titles
        ws.merge_cells("A1:G1")
        ws["A1"] = f"Evaluation AI - Mark Statement: {exam_name}"
        ws["A1"].font = title_font
        ws["A1"].alignment = center_align

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Subject: {subject} | Class: {class_name}"
        ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="475569")
        ws["A2"].alignment = center_align

        headers = ["Rank", "Register Number", "Student Name", "Marks Obtained", "Maximum Marks", "Percentage (%)", "Status"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border

        # Populate rows
        for row_idx, sub in enumerate(submissions, 5):
            rank_val = sub.get("rank") or (row_idx - 4)
            ws.cell(row=row_idx, column=1, value=rank_val).alignment = center_align
            ws.cell(row=row_idx, column=2, value=sub.get("register_number", "")).alignment = center_align
            ws.cell(row=row_idx, column=3, value=sub.get("student_name", "")).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=row_idx, column=4, value=round(sub.get("total_score", 0), 2)).alignment = center_align
            ws.cell(row=row_idx, column=5, value=round(sub.get("max_score", 20), 2)).alignment = center_align
            ws.cell(row=row_idx, column=6, value=f"{sub.get('percentage', 0):.1f}%").alignment = center_align
            ws.cell(row=row_idx, column=7, value=sub.get("status", "").replace("_", " ")).alignment = center_align

            for c in range(1, 8):
                ws.cell(row=row_idx, column=c).border = border

        # Set Column Widths
        column_widths = [10, 20, 30, 18, 18, 18, 20]
        for i, width in enumerate(column_widths, 1):
            col_letter = openpyxl.utils.get_column_letter(i)
            ws.column_dimensions[col_letter].width = width

        wb.save(output_path)
        return output_path

export_service = ExportService()
